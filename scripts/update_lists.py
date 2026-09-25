"""
Ubah daftar Excel di folder data/ menjadi file JS yang dibaca generator.

Jalankan setiap kali Daftar_Saham_IDX.xlsx atau RD_kisi.xlsx diperbarui:
    pip install openpyxl
    python scripts/update_lists.py
Lalu upload ulang data/stocks.js dan data/rd_kisi.js ke GitHub.
(Workflow "Update daftar saham & RD" juga menjalankan ini otomatis saat Excel di-upload.)
"""
import json, re
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"


def short_name(n):
    n = str(n or "").strip()
    n = re.sub(r"\s*Tbk\.?\s*$", "", n)
    n = re.sub(r"\s*\(Persero\)", "", n)
    n = re.sub(r"^PT\.?\s+", "", n)
    return n.strip()


def header_index(row, *names):
    low = [str(c or "").strip().lower() for c in row]
    for n in names:
        if n in low:
            return low.index(n)
    raise SystemExit(f"Kolom {names} tidak ditemukan. Header: {row}")


def stocks():
    ws = openpyxl.load_workbook(DATA / "Daftar_Saham_IDX.xlsx", read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    ic = header_index(rows[0], "kode saham", "kode", "code")
    iname = header_index(rows[0], "nama perusahaan", "nama", "name")
    out = {}
    for r in rows[1:]:
        code = str(r[ic] or "").strip().upper()
        if code:
            out[code] = short_name(r[iname])
    (DATA / "stocks.js").write_text(
        "// Dibuat otomatis oleh scripts/update_lists.py dari Daftar_Saham_IDX.xlsx — jangan edit manual.\n"
        f"window.STOCKS={json.dumps(out, ensure_ascii=False, separators=(',', ':'))};\n", encoding="utf-8")
    print("stocks.js:", len(out), "saham")


def rd():
    ws = openpyxl.load_workbook(DATA / "RD_kisi.xlsx", read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    iname = header_index(rows[0], "name", "nama")
    itype = header_index(rows[0], "type", "jenis", "kategori")
    imin = header_index(rows[0], "min_subs", "minimal pembelian", "min pembelian")
    valid = {"Pasar Uang", "Pendapatan Tetap", "Campuran", "Saham", "Indeks"}
    out = []
    for r in rows[1:]:
        name = str(r[iname] or "").strip()
        if not name:
            continue
        t = str(r[itype] or "").strip()
        if t not in valid:
            raise SystemExit(f"Type tidak dikenal untuk '{name}': '{t}'. Pilihan: {sorted(valid)}")
        out.append([name, t, int(r[imin] or 0)])
    (DATA / "rd_kisi.js").write_text(
        "// Dibuat otomatis oleh scripts/update_lists.py dari RD_kisi.xlsx — jangan edit manual.\n"
        f"window.RD_KISI={json.dumps(out, ensure_ascii=False, separators=(',', ':'))};\n", encoding="utf-8")
    print("rd_kisi.js:", len(out), "produk")


if __name__ == "__main__":
    stocks()
    rd()
